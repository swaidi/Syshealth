"""Duplicate file detection by name or by size within the same extension, with optional content verification."""
from __future__ import annotations

import datetime as dt
import hashlib
import os
import re
import stat
import threading
from functools import lru_cache
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Callable, Optional

from .safety import is_protected_dir
from .platform_utils import IS_WINDOWS, Cancelled, is_protected, is_virtual_dir, list_dir, human_size, norm, \
    protected_roots

FILE_ATTRIBUTE_HIDDEN = 0x2
FILE_ATTRIBUTE_SYSTEM = 0x4
FILE_ATTRIBUTE_REPARSE_POINT = 0x400

SKIP_DIR_NAMES = {"$recycle.bin", "system volume information", "$winreagent", "config.msi",
                  "recovery", "lost+found", ".trash", ".trash-1000", "__pycache__"}
APPDATA_DIR_NAMES = {"appdata", ".cache", ".local", ".config", ".var"}

# Folders that belong to programs or development tools. Their files look like duplicates (every project keeps
# its own copy of the same libraries) but removing one breaks that program. Skipped unless asked for.
PROGRAM_DIR_NAMES = {"site-packages", "dist-packages", "node_modules", "bower_components", ".git", ".svn", ".hg",
                     "__pycache__", ".tox", ".nox", ".mypy_cache", ".pytest_cache", ".ruff_cache", ".gradle",
                     ".m2", ".cargo", ".rustup", ".npm", ".nuget", ".conda", "anaconda3", "miniconda3", ".pyenv",
                     ".nvm", ".wine", "flatpak", ".steam", "steamapps", ".vscode-server", ".cache"}
# Program-library and compiled files: never personal documents
PROGRAM_EXTS = {".dll", ".so", ".dylib", ".pyd", ".pyc", ".pyo", ".o", ".a", ".lib", ".sys", ".ko", ".node",
                ".mui", ".cat"}
_SO_VERSIONED = re.compile(r"\.so(\.\d+)+$", re.I)          # libcufft.so.11, libc.so.6

# System snapshots and backups (Timeshift, Snapper). They hold copies of system files: never duplicates to remove.
SNAPSHOT_DIR_NAMES = {"timeshift", "timeshift-btrfs", ".snapshots", "@snapshots"}
# A folder holding these sub-folders is a copy of a whole system (a snapshot or backup of "/")
SYSTEM_TREE_MARKERS = ("etc", "usr", "bin", "lib", "sbin", "boot", "var")
# A folder named "etc" holding these files is a copy of system settings (for example ~/backup/etc)
ETC_MARKERS = ("fstab", "passwd", "hostname", "os-release", "group", "shadow")
# Audio and video project folders (Ardour, Cubase, Nuendo and similar): their media must stay inside the project
MEDIA_PROJECT_MARKERS = ("interchange", "peaks", "audiofiles")
# Folders whose copies are disposable: when choosing which copy to keep, prefer one outside these
DISPOSABLE_DIR_NAMES = {"log", "logs", "tmp", "temp", "cache", "caches", "build", "dist", "out", "target",
                        "__pycache__", ".cache"}
# File names that mark a copy: "report - Copy.docx", "report (2).docx", "Copy of report.docx", "report_copy.docx"
_COPY_MARK = re.compile(r"^copy of |\s*-\s*copy(\s*\(\d+\))?$|[ _-]copy(\d+)?$|\s*\(\d+\)$|\s+\d+$", re.I)

PARTIAL_BYTES = 64 * 1024
CHUNK = 1024 * 1024


@dataclass
class FileInfo:
    path: str
    name: str
    ext: str
    size: int
    mtime: float
    inode: tuple = ()   # (device, inode) on Linux; the same pair means hard links to one file, not copies


@dataclass
class DupGroup:
    label: str
    files: list[FileInfo] = field(default_factory=list)
    verified: bool = False
    locked: dict = field(default_factory=dict)   # path -> reason this copy must not be deleted

    @property
    def keep(self) -> Optional[FileInfo]:
        return choose_keep(self.files, self.locked)

    @property
    def total(self) -> int:
        return sum(f.size for f in self.files)

    @property
    def wasted(self) -> int:
        """Space freed by deleting every copy that may be deleted except the one to keep."""
        keep = self.keep
        return sum(f.size for f in self.files if f is not keep and f.path not in self.locked)


def normalize_roots(roots: list[str]) -> list[str]:
    """Drop roots nested inside other selected roots so no file is indexed twice."""
    cleaned = sorted({norm(r) for r in roots if os.path.isdir(r)}, key=len)
    result: list[str] = []
    for r in cleaned:
        if not any(r == p or r.startswith(p.rstrip(os.sep) + os.sep) for p in result):
            result.append(r)
    return result


def is_program_dir(path: str, name: str) -> bool:
    """A folder owned by a program or development tool (virtual environments included)."""
    low = name.lower()
    if low in PROGRAM_DIR_NAMES:
        return True
    parent = os.path.dirname(os.path.abspath(path))
    if low == "trash" and os.path.basename(parent) == "share" and \
            os.path.basename(os.path.dirname(parent)) == ".local":
        return True  # Linux Trash: empty it instead of deduplicating it
    try:  # Python virtual environment or conda environment, whatever its name (.venv, env, myenv...)
        if os.path.isfile(os.path.join(path, "pyvenv.cfg")) or os.path.isdir(os.path.join(path, "conda-meta")):
            return True
        # audio or video project (its media files must stay inside the project)
        return any(os.path.isdir(os.path.join(path, m)) for m in MEDIA_PROJECT_MARKERS)
    except OSError:
        return False


@lru_cache(maxsize=50000)
def system_dir_reason(path: str, name: str) -> Optional[str]:
    """Why a folder holds system files (a system folder, a snapshot, or a copy of system files), or None."""
    if os.path.dirname(norm(path)) == norm(path):
        return None  # the real root of a drive
    low = name.lower()
    if low in SNAPSHOT_DIR_NAMES:
        return "system snapshot (Timeshift or similar)"
    if IS_WINDOWS:
        return None
    try:
        if low == "etc" and sum(os.path.exists(os.path.join(path, m)) for m in ETC_MARKERS) >= 2:
            return "copy of system settings"
        if os.path.isdir(os.path.join(path, "etc")) and \
                sum(os.path.isdir(os.path.join(path, m)) for m in SYSTEM_TREE_MARKERS) >= 4:
            return "copy of a system (snapshot or backup)"
    except OSError:
        pass
    return None


def system_reason(path: str) -> Optional[str]:
    """Why a file belongs to the operating system or a copy of it, or None."""
    if is_protected(path):
        return "system folder"
    folder = os.path.dirname(os.path.abspath(path))
    for _ in range(40):
        parent = os.path.dirname(folder)
        if parent == folder:
            break
        reason = system_dir_reason(folder, os.path.basename(folder))
        if reason:
            return reason
        folder = parent
    return None


def is_program_file(name: str) -> bool:
    low = name.lower()
    return os.path.splitext(low)[1] in PROGRAM_EXTS or bool(_SO_VERSIONED.search(low))


def program_reason(path: str) -> Optional[str]:
    """Why a file must stay where it is: system file, snapshot, program, project or library; None if it may go."""
    reason = system_reason(path)
    if reason:
        return reason
    folder = os.path.dirname(os.path.abspath(path))
    for _ in range(12):
        parent = os.path.dirname(folder)
        if parent == folder:
            break
        if is_program_dir(folder, os.path.basename(folder)):
            return "program or project folder"
        folder = parent
    return "program library" if is_program_file(os.path.basename(path)) else None


def in_program_folder(path: str) -> bool:
    """True when a file sits inside a system, program or project folder (used to block or warn before deleting)."""
    return program_reason(path) is not None


def _copy_base(name: str) -> str:
    stem, ext = os.path.splitext(name.lower())
    prev = None
    while prev != stem:
        prev, stem = stem, _COPY_MARK.sub("", stem).strip()
    return stem + ext


def _in_disposable_dir(path: str, common: str = "") -> bool:
    """True when the file sits in a log, temp or build folder below the folder all copies share."""
    folder = norm(os.path.dirname(path))
    if common and (folder == common or folder.startswith(common.rstrip(os.sep) + os.sep)):
        folder = folder[len(common):]
    return any(p.lower() in DISPOSABLE_DIR_NAMES for p in folder.split(os.sep))


def choose_keep(files: list[FileInfo], locked: Optional[dict] = None, newest: bool = True) -> Optional[FileInfo]:
    """The copy to keep: never a locked copy (a protected file may disappear with its program); preferably one
    outside log, temp and build folders, then one with the original name (not "- Copy" or "(2)"), then the
    newest (or the oldest)."""
    free = [f for f in files if not (locked and f.path in locked)]
    if not free:
        return None
    try:
        common = norm(os.path.commonpath([os.path.dirname(f.path) for f in free]))
    except ValueError:  # different drives
        common = ""
    sign = 1 if newest else -1
    return max(free, key=lambda f: (not _in_disposable_dir(f.path, common),
                                    _copy_base(f.name) == f.name.lower(), sign * f.mtime))


def _skip_dir(path: str, name: str, include_hidden: bool, include_appdata: bool, protected: set[str],
              include_program: bool = False) -> bool:
    low = name.lower()
    if low in SKIP_DIR_NAMES or is_protected_dir(path):  # app data is never offered for deletion
        return True
    if system_dir_reason(path, name):  # system snapshots and copies of system files: always skipped
        return True
    if not include_program and is_program_dir(path, name):
        return True
    if not include_appdata and low in APPDATA_DIR_NAMES:
        return True
    if not include_hidden and not IS_WINDOWS and name.startswith("."):
        return True
    return norm(path) in protected


def index_files(roots: list[str], min_size: int = 1, include_hidden: bool = False,
                include_appdata: bool = False, cancel: Optional[threading.Event] = None,
                progress: Optional[Callable[[str], None]] = None,
                include_program: bool = False) -> list[FileInfo]:
    protected = set(protected_roots())
    files: list[FileInfo] = []
    stack = []
    for r in normalize_roots(roots):
        reason = system_reason(os.path.join(r, "x")) or system_dir_reason(r, os.path.basename(r))
        if reason:
            if progress:
                progress(f"Skipped {r}: {reason}. System files are never offered for deletion.")
        else:
            stack.append(r)
    count = 0
    while stack:
        if cancel is not None and cancel.is_set():
            raise Cancelled()
        current = stack.pop()
        entries, _complete = list_dir(current)
        for entry in entries:
            try:
                if entry.is_symlink():
                    continue
                if IS_WINDOWS:
                    st = entry.stat(follow_symlinks=False)
                    attrs = getattr(st, "st_file_attributes", 0)
                    if attrs & FILE_ATTRIBUTE_REPARSE_POINT:
                        continue
                    if not include_hidden and attrs & (FILE_ATTRIBUTE_HIDDEN | FILE_ATTRIBUTE_SYSTEM):
                        continue
                if entry.is_dir(follow_symlinks=False):
                    if not _skip_dir(entry.path, entry.name, include_hidden, include_appdata, protected,
                                     include_program) and not is_virtual_dir(entry.path):
                        stack.append(entry.path)
                    continue
                if not entry.is_file(follow_symlinks=False):
                    continue
                if not include_hidden and not IS_WINDOWS and entry.name.startswith("."):
                    continue
                if not include_program and is_program_file(entry.name):
                    continue
                st = entry.stat(follow_symlinks=False)
                if not stat.S_ISREG(st.st_mode) or st.st_size < min_size:
                    continue
                ext = ".so" if _SO_VERSIONED.search(entry.name) else os.path.splitext(entry.name)[1].lower()
                inode = (st.st_dev, st.st_ino) if st.st_ino and not IS_WINDOWS else ()
                files.append(FileInfo(entry.path, entry.name, ext, st.st_size, st.st_mtime, inode))
                count += 1
                if progress and count % 5000 == 0:
                    progress(f"Indexed {count:,} files… {current}")
            except OSError:
                continue
    if progress:
        progress(f"Indexed {len(files):,} files.")
    return files


def hash_file(path: str, limit: Optional[int] = None, cancel: Optional[threading.Event] = None) -> Optional[str]:
    h = hashlib.sha256()
    remaining = limit
    try:
        with open(path, "rb") as fh:
            while True:
                if cancel is not None and cancel.is_set():
                    raise Cancelled()
                size = CHUNK if remaining is None else min(CHUNK, remaining)
                if size <= 0:
                    break
                data = fh.read(size)
                if not data:
                    break
                h.update(data)
                if remaining is not None:
                    remaining -= len(data)
    except OSError:
        return None
    return h.hexdigest()


def _split_by_content(files: list[FileInfo], cancel, progress) -> list[list[FileInfo]]:
    by_size: dict[int, list[FileInfo]] = defaultdict(list)
    for f in files:
        by_size[f.size].append(f)
    out: list[list[FileInfo]] = []
    for size, same_size in by_size.items():
        if len(same_size) < 2:
            continue
        by_partial: dict[str, list[FileInfo]] = defaultdict(list)
        for f in same_size:
            digest = hash_file(f.path, PARTIAL_BYTES, cancel)
            if digest:
                by_partial[digest].append(f)
        for candidates in by_partial.values():
            if len(candidates) < 2:
                continue
            if size <= PARTIAL_BYTES:  # partial hash already covers the whole file
                out.append(candidates)
                continue
            by_full: dict[str, list[FileInfo]] = defaultdict(list)
            for f in candidates:
                digest = hash_file(f.path, None, cancel)
                if digest:
                    by_full[digest].append(f)
            out.extend(g for g in by_full.values() if len(g) > 1)
    return out


def find_duplicates(files: list[FileInfo], mode: str = "size_ext", verify: bool = True,
                    cancel: Optional[threading.Event] = None,
                    progress: Optional[Callable[[str], None]] = None, show_locked: bool = False) -> list[DupGroup]:
    """mode: 'name' (same file name) or 'size_ext' (same size and same extension).
    show_locked keeps groups in which no copy may be deleted (for reviewing space used by program folders)."""
    buckets: dict[tuple, list[FileInfo]] = defaultdict(list)
    for f in files:
        if mode == "name":
            buckets[(f.name.lower(),)].append(f)
        elif f.size > 0:
            buckets[(f.ext, f.size)].append(f)
    candidates = [b for b in buckets.values() if len(b) > 1]
    if progress:
        progress(f"{len(candidates):,} candidate groups found.")

    groups: list[DupGroup] = []
    total = len(candidates)
    for i, bucket in enumerate(candidates, 1):
        if cancel is not None and cancel.is_set():
            raise Cancelled()
        if verify:
            if progress and i % 200 == 0:
                progress(f"Verifying content: group {i:,} of {total:,}…")
            subgroups = _split_by_content(bucket, cancel, progress)
        else:
            subgroups = [bucket]
        for sub in subgroups:
            sub = _drop_hard_links(sub)
            if len(sub) < 2:
                continue
            locked = _locked_copies(sub, verify)
            if not show_locked and sum(f.path not in locked for f in sub) < 2:
                continue  # nothing in this group may be deleted
            first = sub[0]
            if mode == "name":
                label = first.name
            else:
                label = f"{first.ext or '(no extension)'} · {human_size(first.size)}"
            if verify:
                label += " · identical content"
            groups.append(DupGroup(label, sorted(sub, key=lambda f: f.mtime, reverse=True), verify, locked))
    groups.sort(key=lambda g: g.wasted, reverse=True)
    if progress:
        progress(f"{len(groups):,} duplicate groups; {human_size(sum(g.wasted for g in groups))} reclaimable.")
    return groups


def _drop_hard_links(files: list[FileInfo]) -> list[FileInfo]:
    """Hard links are one file under several names: deleting one frees nothing. Keep one entry per file."""
    seen, out = set(), []
    for f in files:
        if f.inode and f.inode in seen:
            continue
        if f.inode:
            seen.add(f.inode)
        out.append(f)
    return out


def _locked_copies(files: list[FileInfo], verified: bool) -> dict:
    """Copies that must not be deleted, with the reason: system or program files, and files in the same folder
    with different names and identical content (deliberate variants such as firmware_a.bin / firmware_b.bin),
    unless the names mark an ordinary copy ("report - Copy.docx", "report (2).docx")."""
    locked = {}
    for f in files:
        reason = program_reason(f.path)
        if reason:
            locked[f.path] = reason
    if verified:
        by_folder: dict[str, list[FileInfo]] = defaultdict(list)
        for f in files:
            by_folder[norm(os.path.dirname(f.path))].append(f)
        for same in by_folder.values():
            if len(same) > 1 and len({_copy_base(f.name) for f in same}) > 1:
                for f in same:
                    locked.setdefault(f.path, "same folder, different name (deliberate copy)")
    return locked


# --------------------------------------------------------------------------- Excel export
def export_xlsx(groups: list[DupGroup], path: str, locations: list[str], mode: str, verified: bool,
                selected: Optional[set[str]] = None) -> int:
    """Write the duplicate groups to an Excel workbook for review. Returns the number of file rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    selected = {norm(p) for p in (selected or set())}
    wb = Workbook()
    ws = wb.active
    ws.title = "Duplicates"
    head_fill = PatternFill("solid", fgColor="2E3440")
    head_font = Font(bold=True, color="FFFFFF")
    band = PatternFill("solid", fgColor="EEF2F7")
    keep_font = Font(bold=True, color="1E7B34")
    warn_font = Font(color="B3261E")
    headers = ["Group", "Copies", "File name", "Type", "Size (MB)", "Modified", "Folder", "Full path",
               "Suggestion", "Program folder", "Selected in SysHealth", "Your decision"]
    widths = [8, 8, 36, 10, 11, 18, 60, 80, 14, 48, 20, 18]
    ws.append(headers)
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "C2"
    rows = 0
    for gi, g in enumerate(groups, 1):
        keep = g.keep
        for f in sorted(g.files, key=lambda f: f.mtime, reverse=True):
            reason = g.locked.get(f.path)
            program = reason is not None
            suggestion = "Protected" if program else ("Keep" if f is keep else "Duplicate")
            ws.append([gi, len(g.files), f.name, f.ext or "", round(f.size / 1_048_576, 2),
                       dt.datetime.fromtimestamp(f.mtime), os.path.dirname(f.path), f.path, suggestion,
                       f"Yes - do not delete ({reason})" if program else "No",
                       "Yes" if norm(f.path) in selected else "", ""])
            r = ws.max_row
            ws.cell(row=r, column=6).number_format = "yyyy-mm-dd hh:mm"
            ws.cell(row=r, column=5).number_format = "#,##0.00"
            if gi % 2 == 0:
                for col in range(1, len(headers) + 1):
                    ws.cell(row=r, column=col).fill = band
            if f is keep:
                ws.cell(row=r, column=9).font = keep_font
            if program:
                ws.cell(row=r, column=10).font = warn_font
            rows += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"

    info = wb.create_sheet("Summary", 0)
    wasted = sum(g.wasted for g in groups)
    lines = [("SysHealth duplicate report", ""),
             ("Created", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
             ("Locations", "; ".join(locations)),
             ("Matched by", "File name" if mode == "name" else "Size within the same file type"),
             ("Identical content confirmed", "Yes" if verified else "No (same size or name only)"),
             ("Duplicate groups", len(groups)),
             ("Files", rows),
             ("Reclaimable space", human_size(wasted)),
             ("", ""),
             ("How to use this report", ""),
             ("1.", "Review the Duplicates sheet. Filter or sort by any column."),
             ("2.", "'Keep' marks the suggested copy to keep: the newest copy outside log, temp and build folders."),
             ("3.", "'Program folder: Yes' (suggestion 'Protected') means a system, snapshot, program or project "
                    "file, or a deliberate copy in the same folder. SysHealth does not let you delete it."),
             ("4.", "Use 'Your decision' to note what to keep or delete, then act on it in SysHealth."),
             ("5.", "Nothing is deleted by this report. Deleted files go to the Recycle Bin / Trash.")]
    for k, v in lines:
        info.append([k, v])
    info["A1"].font = Font(bold=True, size=14)
    info["A10"].font = Font(bold=True)
    for r in range(2, 9):
        info.cell(row=r, column=1).font = Font(bold=True)
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 90
    wb.save(path)
    return rows
